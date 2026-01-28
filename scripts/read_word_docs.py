"""
读取Word文档内容并转换为文本
"""
import os
from docx import Document

# 读取Word文档的函数
def read_docx(file_path):
    """读取Word文档内容"""
    doc = Document(file_path)
    content = []
    
    # 读取段落
    for para in doc.paragraphs:
        if para.text.strip():
            content.append(para.text.strip())
    
    # 读取表格
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                if cell.text.strip():
                    row_text.append(cell.text.strip())
            if row_text:
                content.append(" | ".join(row_text))
    
    return "\n\n".join(content)

# 读取所有Word文档
docs_dir = "assets"
doc_files = [
    "成功CEO临门一脚   数据合规管理1.docx",
    "成功CEO临门一脚  数据合规管理2.docx",
    "成功CEO临门一脚 数据合规管理3.docx",
    "成功CEO临门一脚 数据合规管理4.docx"
]

all_content = []
for doc_file in doc_files:
    doc_path = os.path.join(docs_dir, doc_file)
    print(f"正在读取: {doc_file}")
    
    try:
        content = read_docx(doc_path)
        all_content.append(f"## {doc_file}\n\n{content}")
        print(f"  ✓ 读取成功，内容长度: {len(content)} 字符")
    except Exception as e:
        print(f"  ✗ 读取失败: {str(e)}")

# 合并所有内容并保存
output_file = "assets/数据合规管理知识库.md"
with open(output_file, 'w', encoding='utf-8') as f:
    f.write("# 数据合规管理知识库\n\n")
    f.write("\n\n".join(all_content))

print(f"\n所有文档已合并保存到: {output_file}")
print(f"总内容长度: {len(open(output_file, 'r', encoding='utf-8').read())} 字符")
