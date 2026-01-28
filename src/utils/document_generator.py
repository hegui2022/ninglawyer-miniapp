"""
合同文档生成器
Contract Document Generator
支持 PDF 和 Word 格式导出
"""

from typing import Dict, Any
import os
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.units import cm


class ContractDocumentGenerator:
    """合同文档生成器"""
    
    def __init__(self, output_dir: str = "/tmp"):
        """
        初始化文档生成器
        
        Args:
            output_dir: 输出目录
        """
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)
    
    def generate_word(self, contract_data: Dict[str, Any]) -> str:
        """
        生成 Word 格式合同
        
        Args:
            contract_data: 合同数据
            
        Returns:
            生成的文件路径
        """
        # 创建 Word 文档
        doc = Document()
        
        # 设置文档样式
        self._setup_word_styles(doc)
        
        # 添加合同标题
        self._add_word_title(doc, contract_data)
        
        # 添加合同内容
        self._add_word_content(doc, contract_data)
        
        # 保存文件
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"contract_{contract_data.get('id', 'new')}_{timestamp}.docx"
        filepath = os.path.join(self.output_dir, filename)
        doc.save(filepath)
        
        return filepath
    
    def generate_pdf(self, contract_data: Dict[str, Any]) -> str:
        """
        生成 PDF 格式合同
        
        Args:
            contract_data: 合同数据
            
        Returns:
            生成的文件路径
        """
        # 创建 PDF 文档
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"contract_{contract_data.get('id', 'new')}_{timestamp}.pdf"
        filepath = os.path.join(self.output_dir, filename)
        
        # 创建 PDF
        story = []
        doc = SimpleDocTemplate(
            filepath,
            pagesize=A4,
            rightMargin=2*cm,
            leftMargin=2*cm,
            topMargin=2*cm,
            bottomMargin=2*cm
        )
        
        # 添加标题
        story.extend(self._get_pdf_title(contract_data))
        
        # 添加内容
        story.extend(self._get_pdf_content(contract_data))
        
        # 构建文档
        doc.build(story)
        
        return filepath
    
    def _setup_word_styles(self, doc: Document):
        """设置 Word 文档样式"""
        # 设置默认字体
        style = doc.styles['Normal']
        font = style.font
        font.name = '宋体'
        font.size = Pt(12)
        
        # 标题样式
        title_style = doc.styles['Heading 1']
        title_font = title_style.font
        title_font.name = '黑体'
        title_font.size = Pt(16)
        title_font.bold = True
        title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    def _add_word_title(self, doc: Document, contract_data: Dict[str, Any]):
        """添加 Word 标题"""
        # 合同标题
        title = doc.add_heading()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(self._get_contract_title(contract_data))
        run.font.name = '黑体'
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        
        # 空行
        doc.add_paragraph()
    
    def _add_word_content(self, doc: Document, contract_data: Dict[str, Any]):
        """添加 Word 内容"""
        content = contract_data.get('contract_content', {})
        
        # 双方信息
        doc.add_heading('合同双方', level=1)
        p = doc.add_paragraph()
        p.add_run(f"甲方（用人单位）：{content.get('party_a', '')}")
        p.add_run(f"\n乙方（劳动者）：{content.get('party_b', '')}")
        
        # 合同期限
        if content.get('start_date') or content.get('end_date'):
            doc.add_heading('合同期限', level=1)
            p = doc.add_paragraph()
            p.add_run(f"本合同期限自 {content.get('start_date', '')} 至 {content.get('end_date', '')}")
            if content.get('probation_months'):
                p.add_run(f"\n试用期：{content.get('probation_months')} 个月")
        
        # 工作内容
        if content.get('job_content'):
            doc.add_heading('工作内容', level=1)
            p = doc.add_paragraph(content.get('job_content'))
        
        # 工作地点
        if content.get('workplace'):
            doc.add_heading('工作地点', level=1)
            p = doc.add_paragraph(content.get('workplace'))
        
        # 劳动报酬
        if content.get('salary'):
            doc.add_heading('劳动报酬', level=1)
            p = doc.add_paragraph()
            p.add_run(f"月薪：{content.get('salary')} 元")
            if content.get('payment_method'):
                p.add_run(f"\n支付方式：{content.get('payment_method')}")
        
        # 工作时间
        if content.get('work_hours'):
            doc.add_heading('工作时间', level=1)
            p = doc.add_paragraph(content.get('work_hours'))
        
        # 社会保险
        if content.get('has_insurance') is not None:
            doc.add_heading('社会保险', level=1)
            p = doc.add_paragraph("缴纳" if content.get('has_insurance') else "不缴纳")
        
        # 住房公积金
        if content.get('has_housing_fund') is not None:
            doc.add_heading('住房公积金', level=1)
            p = doc.add_paragraph("缴纳" if content.get('has_housing_fund') else "不缴纳")
        
        # 附加条款
        clauses = []
        if content.get('has_confidentiality'):
            clauses.append("保密协议")
        if content.get('has_non_compete'):
            clauses.append("竞业限制")
        if content.get('has_ip_clause'):
            clauses.append("知识产权条款")
        
        if clauses:
            doc.add_heading('附加条款', level=1)
            for clause in clauses:
                doc.add_paragraph(f"• {clause}", style='List Bullet')
        
        # 自定义条款
        if content.get('custom_clause'):
            doc.add_heading('自定义条款', level=1)
            p = doc.add_paragraph(content.get('custom_clause'))
        
        # 签署信息
        doc.add_page_break()
        doc.add_heading('签署', level=1)
        
        table = doc.add_table(rows=3, cols=2)
        table.style = 'Table Grid'
        
        # 甲方签署
        table.rows[0].cells[0].text = '甲方（用人单位）：'
        table.rows[0].cells[1].text = '乙方（劳动者）：'
        table.rows[1].cells[0].text = '法定代表人/授权代表：'
        table.rows[1].cells[1].text = '身份证号码：'
        table.rows[2].cells[0].text = '签署日期：'
        table.rows[2].cells[1].text = '签署日期：'
    
    def _get_pdf_title(self, contract_data: Dict[str, Any]):
        """获取 PDF 标题"""
        styles = getSampleStyleSheet()
        title_style = styles['Heading1']
        title_style.alignment = 1  # 居中
        title_style.fontSize = 20
        title_style.leading = 30
        
        title = Paragraph(self._get_contract_title(contract_data), title_style)
        spacer = Spacer(1, 0.5*cm)
        
        return [title, spacer]
    
    def _get_pdf_content(self, contract_data: Dict[str, Any]):
        """获取 PDF 内容"""
        styles = getSampleStyleSheet()
        normal_style = styles['Normal']
        heading_style = styles['Heading2']
        
        content = []
        content_obj = contract_data.get('contract_content', {})
        
        # 合同双方
        content.append(Paragraph('合同双方', heading_style))
        content.append(Paragraph(f"甲方（用人单位）：{content_obj.get('party_a', '')}", normal_style))
        content.append(Paragraph(f"乙方（劳动者）：{content_obj.get('party_b', '')}", normal_style))
        content.append(Spacer(1, 0.3*cm))
        
        # 合同期限
        if content_obj.get('start_date') or content_obj.get('end_date'):
            content.append(Paragraph('合同期限', heading_style))
            term_text = f"本合同期限自 {content_obj.get('start_date', '')} 至 {content_obj.get('end_date', '')}"
            if content_obj.get('probation_months'):
                term_text += f"<br/>试用期：{content_obj.get('probation_months')} 个月"
            content.append(Paragraph(term_text, normal_style))
            content.append(Spacer(1, 0.3*cm))
        
        # 工作内容
        if content_obj.get('job_content'):
            content.append(Paragraph('工作内容', heading_style))
            content.append(Paragraph(content_obj.get('job_content'), normal_style))
            content.append(Spacer(1, 0.3*cm))
        
        # 工作地点
        if content_obj.get('workplace'):
            content.append(Paragraph('工作地点', heading_style))
            content.append(Paragraph(content_obj.get('workplace'), normal_style))
            content.append(Spacer(1, 0.3*cm))
        
        # 劳动报酬
        if content_obj.get('salary'):
            content.append(Paragraph('劳动报酬', heading_style))
            salary_text = f"月薪：{content_obj.get('salary')} 元"
            if content_obj.get('payment_method'):
                salary_text += f"<br/>支付方式：{content_obj.get('payment_method')}"
            content.append(Paragraph(salary_text, normal_style))
            content.append(Spacer(1, 0.3*cm))
        
        # 工作时间
        if content_obj.get('work_hours'):
            content.append(Paragraph('工作时间', heading_style))
            content.append(Paragraph(content_obj.get('work_hours'), normal_style))
            content.append(Spacer(1, 0.3*cm))
        
        # 社会保险
        if content_obj.get('has_insurance') is not None:
            content.append(Paragraph('社会保险', heading_style))
            content.append(Paragraph("缴纳" if content_obj.get('has_insurance') else "不缴纳", normal_style))
            content.append(Spacer(1, 0.3*cm))
        
        # 住房公积金
        if content_obj.get('has_housing_fund') is not None:
            content.append(Paragraph('住房公积金', heading_style))
            content.append(Paragraph("缴纳" if content_obj.get('has_housing_fund') else "不缴纳", normal_style))
            content.append(Spacer(1, 0.3*cm))
        
        # 附加条款
        clauses = []
        if content_obj.get('has_confidentiality'):
            clauses.append("保密协议")
        if content_obj.get('has_non_compete'):
            clauses.append("竞业限制")
        if content_obj.get('has_ip_clause'):
            clauses.append("知识产权条款")
        
        if clauses:
            content.append(Paragraph('附加条款', heading_style))
            for clause in clauses:
                content.append(Paragraph(f"• {clause}", normal_style))
            content.append(Spacer(1, 0.3*cm))
        
        # 自定义条款
        if content_obj.get('custom_clause'):
            content.append(Paragraph('自定义条款', heading_style))
            content.append(Paragraph(content_obj.get('custom_clause'), normal_style))
            content.append(Spacer(1, 0.3*cm))
        
        # 签署信息
        content.append(Spacer(1, 0.5*cm))
        content.append(Paragraph('签署', heading_style))
        
        sign_data = [
            ['甲方（用人单位）：', '乙方（劳动者）：'],
            ['法定代表人/授权代表：', '身份证号码：'],
            ['签署日期：', '签署日期：']
        ]
        
        sign_table = Table(sign_data, colWidths=[8*cm, 8*cm])
        sign_table.setStyle(TableStyle([
            ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 12),
            ('RIGHTPADDING', (0, 0), (-1, -1), 12),
            ('TOPPADDING', (0, 0), (-1, -1), 12),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
        ]))
        content.append(sign_table)
        
        return content
    
    def _get_contract_title(self, contract_data: Dict[str, Any]) -> str:
        """获取合同标题"""
        contract_type = contract_data.get('contract_type', '')
        name = contract_data.get('contract_name', '')
        
        type_map = {
            'standard': '劳动合同',
            'parttime': '非全日制用工合同',
            'intern': '实习协议',
            'retired': '退休返聘协议',
            'project': '项目制合同',
            'dispatch': '劳务派遣合同'
        }
        
        if name:
            return name
        return type_map.get(contract_type, '合同')


# 全局实例
document_generator = ContractDocumentGenerator()
